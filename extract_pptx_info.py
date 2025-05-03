from pptx import Presentation
import os

# Path to the PowerPoint file
pptx_path = "test.pptx"

# Check if file exists
if not os.path.exists(pptx_path):
    print(f"Error: File '{pptx_path}' not found.")
    exit(1)

try:
    # Load presentation
    prs = Presentation(pptx_path)
    
    # Print summary
    print(f"\nAnalyzing PowerPoint: {pptx_path}")
    print(f"Total slides: {len(prs.slides)}\n")
    
    # Analyze each slide
    for i, slide in enumerate(prs.slides, 1):
        # Try to get slide title
        title = "No title found"
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.has_text_frame:
                text = shape.text.strip()
                if text and (shape.name.startswith("Title") or "title" in shape.name.lower() or i == 1 and text):
                    title = text
                    break
        
        # Get slide text content (snippets for context)
        content = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.has_text_frame:
                text = shape.text.strip()
                if text and text != title:
                    # Add first few words as preview
                    words = text.split()
                    preview = " ".join(words[:5])
                    if len(words) > 5:
                        preview += "..."
                    content.append(preview)
        
        # Print slide information
        print(f"Slide {i}:")
        print(f"  Title: {title}")
        if content:
            print(f"  Content preview: {' | '.join(content[:3])}")
        print()

except Exception as e:
    print(f"Error analyzing presentation: {str(e)}") 